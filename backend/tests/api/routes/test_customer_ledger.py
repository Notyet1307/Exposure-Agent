import hashlib
import io
import re
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from typing import Any, NoReturn, cast

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from sqlalchemy import text
from sqlalchemy.exc import SQLAlchemyError
from sqlmodel import Session, select

from app.core.config import settings
from app.core.db import engine
from app.domain import customer_ledger as ledger
from app.domain.models import Artifact, CustomerLedgerRevision, CustomerUpload
from tests.api.routes.test_customer_uploads import _create_member, _create_project
from tests.utils.audit import reject_audit_inserts

LedgerSetup = tuple[str, dict[str, str], dict[str, Any], bytes, Path]


@pytest.fixture
def setup_ledger(
    client: TestClient,
    superuser_token_headers: dict[str, str],
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    request: pytest.FixtureRequest,
) -> LedgerSetup:
    monkeypatch.setattr(settings, "ARTIFACT_ROOT", tmp_path)
    project = _create_project(client, superuser_token_headers)
    root = f"{settings.API_V1_STR}/projects/{project['id']}"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(list(ledger.HEADERS))
    for n in range(getattr(request, "param", 61)):
        # Duplicate IPs are distinct original rows, and optional cell types survive editing.
        sheet.append(
            [
                "2001:db8:10::20",
                443,
                443,
                "是",
                "example.net",
                "HTTPS",
                "原负责人",
                "业务部",
                "端口负责人",
                datetime(2026, 1, 2, 3, 4),
                n + 1,
            ]
        )
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    data = stream.getvalue()
    uploaded = client.post(
        root + "/customer-uploads",
        headers=superuser_token_headers,
        files={"file": ("original.xlsx", data)},
    )
    assert uploaded.status_code == 201, uploaded.text
    upload = uploaded.json()
    assert (
        client.post(
            root + f"/customer-uploads/{upload['id']}/select",
            headers=superuser_token_headers,
        ).status_code
        == 200
    )
    return root, superuser_token_headers, upload, data, tmp_path


def read(client: TestClient, setup: LedgerSetup, **params: Any) -> dict[str, Any]:
    root, headers, *_ = setup
    response = client.get(root + "/customer-ledger", headers=headers, params=params)
    assert response.status_code == 200, response.text
    return cast(dict[str, Any], response.json())


def edit(
    page: dict[str, Any],
    operation: str = "update",
    fields: dict[str, Any] | None = None,
    **extra: Any,
) -> dict[str, Any]:
    return dict(
        expected_upload_id=page["upload_id"],
        expected_revision_id=page["revision_id"],
        expected_profile_id=page["current_profile_id"],
        operation=operation,
        entry_id=page["data"][0]["entry_id"],
        fields=fields or {"asset_ip": "2001:db8:10::25"},
        reason="更正录入地址",
        **extra,
    )


def post(
    client: TestClient,
    setup: LedgerSetup,
    payload: dict[str, Any],
    key: str | None = None,
) -> Any:
    root, headers, *_ = setup
    return client.post(
        root + "/customer-ledger/revisions",
        headers={**headers, "Idempotency-Key": key or str(uuid.uuid4())},
        json=payload,
    )


def test_original_pagination_edit_history_and_typed_fields(
    client: TestClient, setup_ledger: LedgerSetup
) -> None:
    first = read(client, setup_ledger)
    second = read(client, setup_ledger, upload_id=first["upload_id"], skip=25)
    assert (
        first["count"] == 61 and len(first["data"]) == 25 and first["unique_ips"] == 1
    )
    assert first["data"][0]["entry_id"] != second["data"][0]["entry_id"]
    assert first["revision_id"] is None and first["can_edit"]
    changed = post(client, setup_ledger, edit(first))
    assert changed.status_code == 201, changed.text
    current = read(client, setup_ledger)
    assert (
        current["revision_id"] == changed.json()["id"]
        and current["upload_id"] != first["upload_id"]
    )
    assert (
        current["unique_ips"] == 2
        and current["data"][0]["canonical_ip"] == "2001:db8:10::25"
    )
    assert current["data"][1]["canonical_ip"] == "2001:db8:10::20"
    assert current["data"][0]["entry_id"] == first["data"][0]["entry_id"]
    old = read(client, setup_ledger, upload_id=first["upload_id"])
    assert old["data"] == first["data"] and not old["can_edit"]
    assert (
        current["data"][0]["fields"]["department"]
        == first["data"][0]["fields"]["department"]
    )
    with Session(engine) as db:
        u = db.get(CustomerUpload, current["upload_id"])
        assert u is not None
        a = db.get(Artifact, u.artifact_id)
        assert a is not None
        path = setup_ledger[-1] / a.storage_key
        book = load_workbook(path)
        assert book.active is not None
        assert book.active.cell(2, 10).value == datetime(2026, 1, 2, 3, 4)
        book.close()
        original = db.get(CustomerUpload, first["upload_id"])
        assert original is not None
        source = db.get(Artifact, original.artifact_id)
        assert source is not None
        assert (
            hashlib.sha256(
                (setup_ledger[-1] / source.storage_key).read_bytes()
            ).hexdigest()
            == hashlib.sha256(setup_ledger[3]).hexdigest()
        )


def test_replay_conflict_metadata_and_old_selection(
    client: TestClient, setup_ledger: LedgerSetup
) -> None:
    first = read(client, setup_ledger)
    body = edit(first)
    key = str(uuid.uuid4())
    a = post(client, setup_ledger, body, key)
    b = post(client, setup_ledger, body, key)
    assert a.status_code == b.status_code == 201 and a.json() == b.json()
    assert (
        post(client, setup_ledger, {**body, "reason": "其他意图"}, key).status_code
        == 409
    )
    assert post(client, setup_ledger, body).status_code == 409
    current = read(client, setup_ledger)
    meta = edit(current, operation="manage")
    meta["fields"] = {}
    meta["management"] = {"owner": "人工确认", "tags": ["关注"]}
    r = post(client, setup_ledger, meta)
    assert r.status_code == 201, r.text
    newer = read(client, setup_ledger)
    assert (
        newer["upload_id"] == current["upload_id"]
        and newer["revision_id"] != current["revision_id"]
    )
    assert newer["data"][0]["management"]["owner"] == "人工确认"
    assert newer["data"][0]["fields"]["asset_owner"] == "原负责人"
    root, headers, *_ = setup_ledger
    assert (
        client.get(root + f"/customer-ledger/operations/{key}", headers=headers).json()[
            "id"
        ]
        == a.json()["id"]
    )
    assert (
        client.post(
            root + f"/customer-uploads/{first['upload_id']}/select", headers=headers
        ).status_code
        == 200
    )
    selected = read(client, setup_ledger)
    assert (
        selected["revision_id"] is None
        and selected["data"][0]["management"]["owner"] == ""
    )
    assert (
        read(client, setup_ledger, revision_id=r.json()["id"])["data"][0]["management"][
            "owner"
        ]
        == "人工确认"
    )


def test_concurrent_saves_have_one_winner(
    client: TestClient, setup_ledger: LedgerSetup
) -> None:
    first = read(client, setup_ledger)
    body = edit(first)
    with ThreadPoolExecutor(max_workers=2) as workers:
        statuses = list(
            workers.map(
                lambda _: post(client, setup_ledger, body).status_code, range(2)
            )
        )
    assert sorted(statuses) == [201, 409]


def test_archive_add_validation_and_immutability(
    client: TestClient, setup_ledger: LedgerSetup
) -> None:
    first = read(client, setup_ledger)
    body = edit(first, operation="archive")
    body["fields"] = {}
    r = post(client, setup_ledger, body)
    assert r.status_code == 201, r.text
    current = read(client, setup_ledger)
    assert current["count"] == 60
    archived = read(client, setup_ledger, archived=True)
    assert archived["count"] == 1
    body = edit(current, operation="add", fields=first["data"][0]["fields"])
    body["entry_id"] = None
    added = post(client, setup_ledger, body)
    assert added.status_code == 201, added.text
    current = read(client, setup_ledger)
    assert current["count"] == 61
    assert (
        post(
            client, setup_ledger, edit(current, fields={"start_port": 1, "end_port": 2})
        ).status_code
        == 422
    )
    assert (
        post(
            client, setup_ledger, edit(current, fields={"asset_ip": "fe80::1%en0"})
        ).status_code
        == 422
    )
    with Session(engine) as db:
        with pytest.raises(SQLAlchemyError, match="immutable"):
            db.execute(
                text(
                    "UPDATE customer_ledger_revisions SET reason=:reason WHERE id=:id"
                ),
                {"reason": "changed", "id": added.json()["id"]},
            )
            db.commit()
        db.rollback()
        with pytest.raises(SQLAlchemyError, match="immutable"):
            db.execute(
                text(
                    "DELETE FROM customer_ledger_entry_versions WHERE revision_id=:id"
                ),
                {"id": added.json()["id"]},
            )
            db.commit()
        db.rollback()


def test_permissions_recovery_scope_and_archive(
    client: TestClient, setup_ledger: LedgerSetup
) -> None:
    root, admin, *_ = setup_ledger
    project_id = root.split("/")[-1]
    first = read(client, setup_ledger)
    op = post(client, setup_ledger, edit(first), "scope-key")
    assert op.status_code == 201
    for roles in [["viewer"], ["approver"]]:
        headers = _create_member(client, admin, project_id=project_id, roles=roles)
        assert (
            client.get(root + "/customer-ledger", headers=headers).json()["can_edit"]
            is False
        )
        assert (
            client.post(
                root + "/customer-ledger/revisions",
                headers={**headers, "Idempotency-Key": "denied"},
                json=edit(first),
            ).status_code
            == 404
        )
        assert (
            client.get(
                root + "/customer-ledger/operations/scope-key", headers=headers
            ).status_code
            == 404
        )
    other = _create_project(client, admin)
    assert (
        client.get(
            f"{settings.API_V1_STR}/projects/{other['id']}/customer-ledger",
            headers=admin,
            params={"revision_id": op.json()["id"]},
        ).status_code
        == 404
    )
    assert client.post(root + "/archive", headers=admin).status_code == 200
    assert (
        post(client, setup_ledger, edit(read(client, setup_ledger))).status_code == 409
    )


@pytest.mark.parametrize("failure", ["audit", "commit", "writer"])
def test_failure_keeps_input_and_cleans_only_new_artifacts(
    client: TestClient,
    setup_ledger: LedgerSetup,
    monkeypatch: pytest.MonkeyPatch,
    failure: str,
) -> None:
    first = read(client, setup_ledger)
    files = set(setup_ledger[-1].rglob("*.xlsx"))

    def fail(*_args: Any, **_kwargs: Any) -> NoReturn:
        raise SQLAlchemyError("injected")

    if failure == "audit":
        with Session(engine) as db, reject_audit_inserts(db):
            result = post(client, setup_ledger, edit(first))
    else:
        with monkeypatch.context() as patch:
            patch.setattr(
                Session, "commit", fail
            ) if failure == "commit" else patch.setattr(ledger, "_write_workbook", fail)
            result = post(client, setup_ledger, edit(first))
    assert result.status_code == 500, result.text
    assert read(client, setup_ledger)["upload_id"] == first["upload_id"]
    assert set(setup_ledger[-1].rglob("*.xlsx")) == files
    with Session(engine) as db:
        assert not db.exec(
            select(CustomerLedgerRevision).where(
                CustomerLedgerRevision.project_id == uuid.UUID(first["project_id"])
            )
        ).all()


@pytest.mark.parametrize("setup_ledger", [1], indirect=True)
def test_last_entry_and_invalid_management_rejected(
    client: TestClient, setup_ledger: LedgerSetup
) -> None:
    first = read(client, setup_ledger)
    body = edit(first, operation="archive")
    body["fields"] = {}
    response = post(client, setup_ledger, body)
    assert (
        response.status_code == 422
        and response.json()["detail"]["code"] == "ledger_last_entry"
    )
    assert read(client, setup_ledger)["upload_id"] == first["upload_id"]
    body = edit(first, operation="manage")
    body["fields"] = {}
    body["management"] = {"tags": ["x" * 65]}
    assert post(client, setup_ledger, body).status_code == 422


def test_lost_commit_acknowledgement_recovers_committed_artifact(
    client: TestClient, setup_ledger: LedgerSetup, monkeypatch: pytest.MonkeyPatch
) -> None:
    first = read(client, setup_ledger)
    key = "lost-commit"
    original_commit = Session.commit

    def commit_then_fail(session: Session) -> None:
        original_commit(session)
        raise SQLAlchemyError("lost acknowledgement")

    with monkeypatch.context() as patch:
        patch.setattr(Session, "commit", commit_then_fail)
        response = post(client, setup_ledger, edit(first), key)
    assert response.status_code == 201, response.text
    current = read(client, setup_ledger)
    assert current["revision_id"] == response.json()["id"]
    with Session(engine) as session:
        upload = session.get(CustomerUpload, current["upload_id"])
        assert upload is not None
        artifact = session.get(Artifact, upload.artifact_id)
        assert artifact is not None
        assert (setup_ledger[-1] / artifact.storage_key).is_file()


@pytest.mark.parametrize("dimension", ["A1:K2", "A1:XFD1048576"])
def test_declared_dimensions_cannot_truncate_or_expand_ledger(
    client: TestClient, setup_ledger: LedgerSetup, dimension: str
) -> None:
    root, headers, _upload, raw, _files = setup_ledger
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw)) as original,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as replacement,
    ):
        for info in original.infolist():
            content = original.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                content = re.sub(
                    rb'<dimension ref="[^"]+"',
                    f'<dimension ref="{dimension}"'.encode(),
                    content,
                )
            replacement.writestr(info, content)
    accepted = client.post(
        root + "/customer-uploads",
        headers=headers,
        files={"file": ("declared-dimension.xlsx", output.getvalue())},
    )
    assert accepted.status_code == 201, accepted.text
    upload_id = accepted.json()["id"]
    assert (
        client.post(
            root + f"/customer-uploads/{upload_id}/select", headers=headers
        ).status_code
        == 200
    )
    before = read(client, setup_ledger)
    assert before["count"] == 61
    result = post(client, setup_ledger, edit(before))
    assert result.status_code == 201, result.text
    assert read(client, setup_ledger)["count"] == 61


@pytest.mark.parametrize(
    "fields",
    [
        {"asset_owner": "bad\u0001owner"},
        {"department": {"kind": "datetime", "value": "2026-09-15T12:00:00+08:00"}},
    ],
)
def test_unsupported_cell_values_are_stable_validation_failures(
    client: TestClient, setup_ledger: LedgerSetup, fields: dict[str, Any]
) -> None:
    first = read(client, setup_ledger)
    response = post(client, setup_ledger, edit(first, fields=fields))
    assert (
        response.status_code == 422
        and response.json()["detail"]["code"] == "ledger_fields_invalid"
    )
    assert read(client, setup_ledger)["upload_id"] == first["upload_id"]


def test_unchanged_input_does_not_materialize_another_workbook(
    client: TestClient, setup_ledger: LedgerSetup
) -> None:
    first = read(client, setup_ledger)
    result = post(client, setup_ledger, edit(first, fields=first["data"][0]["fields"]))
    assert result.status_code == 201, result.text
    assert result.json()["input_changed"] is False
    assert result.json()["upload_id"] == first["upload_id"]


@pytest.mark.parametrize("bad", ["null\x00text", "bad\ud800text"])
def test_invalid_database_text_is_rejected_before_an_intent(
    client: TestClient, setup_ledger: LedgerSetup, bad: str
) -> None:
    first = read(client, setup_ledger)
    body = edit(first, operation="manage")
    body["fields"] = {}
    body["management"] = {"owner": bad}
    root, headers, *_ = setup_ledger
    import json

    result = client.post(
        root + "/customer-ledger/revisions",
        headers={
            **headers,
            "Idempotency-Key": "bad-text",
            "Content-Type": "application/json",
        },
        content=json.dumps(body, ensure_ascii=True),
    )
    assert result.status_code == 422
    assert read(client, setup_ledger)["revision_id"] == first["revision_id"]
