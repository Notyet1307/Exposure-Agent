"""Versioned customer authoring; GovernanceRun still pins CustomerUpload only."""

from __future__ import annotations

import copy
import hashlib
import json
import tempfile
import uuid
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Annotated, Any, Literal, cast

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import (  # type: ignore[import-untyped]
    IllegalCharacterError,
)
from pydantic import (
    AfterValidator,
    BaseModel,
    ConfigDict,
    Field,
    FiniteFloat,
    StrictBool,
    StrictInt,
)
from sqlalchemy.exc import SQLAlchemyError
from sqlmodel import Session, col, select

from app.core.config import settings
from app.domain import customer_uploads as uploads
from app.domain.customer_upload_profiles import REQUIRED_HEADERS, WARNING_HEADERS
from app.domain.customer_upload_validator import (
    MAX_WORKBOOK_BYTES,
    CustomerUploadValidationError,
    _validate_required_row,
    validate_customer_upload_workbook,
)
from app.domain.ip_consistency import IPRecordContractError, normalize_ip
from app.domain.models import (
    Artifact,
    AuditEvent,
    CustomerLedgerEntryVersion,
    CustomerLedgerRevision,
    CustomerUpload,
    Project,
)

FIELD_NAMES: tuple[FieldName, ...] = (
    "asset_ip",
    "start_port",
    "end_port",
    "is_web",
    "web_url",
    "service_type",
    "asset_owner",
    "asset_department",
    "port_owner",
    "department",
    "serial",
)
HEADERS = (*REQUIRED_HEADERS, *WARNING_HEADERS, "序号")
FieldName = Literal[
    "asset_ip",
    "start_port",
    "end_port",
    "is_web",
    "web_url",
    "service_type",
    "asset_owner",
    "asset_department",
    "port_owner",
    "department",
    "serial",
]


def _valid_text(value: str) -> str:
    if "\x00" in value:
        raise ValueError("unsupported_text")
    try:
        value.encode("utf-8")
    except UnicodeError:
        raise ValueError("unsupported_text") from None
    return value


Text = Annotated[str, AfterValidator(_valid_text)]


class TemporalCell(BaseModel):
    model_config = ConfigDict(extra="forbid")
    kind: Literal["datetime", "date", "time", "timedelta"]
    value: Annotated[Text, Field(max_length=128)]


Cell = (
    Annotated[Text, Field(max_length=32767)]
    | StrictInt
    | FiniteFloat
    | StrictBool
    | None
    | TemporalCell
)


class Management(BaseModel):
    model_config = ConfigDict(extra="forbid")
    owner: Text = Field(default="", max_length=255)
    department: Text = Field(default="", max_length=255)
    tags: list[Annotated[Text, Field(min_length=1, max_length=64)]] = Field(
        default_factory=list, max_length=20
    )
    followed: bool = False


class LedgerEntry(BaseModel):
    model_config = ConfigDict(extra="forbid")
    entry_id: uuid.UUID
    position: int
    canonical_ip: str
    fields: dict[FieldName, Cell]
    management: Management = Field(default_factory=Management)
    archived: bool = False


class LedgerEdit(BaseModel):
    model_config = ConfigDict(extra="forbid")
    expected_upload_id: uuid.UUID
    expected_revision_id: uuid.UUID | None
    expected_profile_id: uuid.UUID
    operation: Literal["add", "update", "archive", "manage"]
    entry_id: uuid.UUID | None = None
    fields: dict[FieldName, Cell] = Field(default_factory=dict)
    management: Management | None = None
    reason: Text = Field(min_length=1, max_length=1000)


class RevisionPublic(BaseModel):
    id: uuid.UUID
    parent_revision_id: uuid.UUID | None
    base_upload_id: uuid.UUID
    upload_id: uuid.UUID
    created_by: uuid.UUID
    created_at: datetime
    reason: str
    input_changed: bool


class LedgerPage(BaseModel):
    project_id: uuid.UUID
    upload_id: uuid.UUID | None
    revision_id: uuid.UUID | None
    current_upload_id: uuid.UUID | None
    current_revision_id: uuid.UUID | None
    current_profile_id: uuid.UUID
    upload_sha256: str | None = None
    filename: str | None = None
    source_created_at: datetime | None = None
    revision: RevisionPublic | None = None
    data: list[LedgerEntry] = Field(default_factory=list)
    count: int = 0
    total_records: int = 0
    unique_ips: int = 0
    can_edit: bool = False


class LedgerError(Exception):
    def __init__(self, code: str, status: int = 409) -> None:
        self.code = code
        self.status = status


def public_revision(record: CustomerLedgerRevision) -> RevisionPublic:
    return RevisionPublic.model_validate(record, from_attributes=True)


def _cell(value: Any) -> Any:
    if isinstance(value, datetime | date | time):
        return {"kind": type(value).__name__, "value": value.isoformat()}
    if isinstance(value, timedelta):
        return {"kind": "timedelta", "value": str(value.total_seconds())}
    return value


def _excel(value: Any) -> Any:
    if isinstance(value, TemporalCell):
        value = value.model_dump()
    if isinstance(value, dict):
        cell = TemporalCell.model_validate(value)
        if cell.kind == "timedelta":
            return timedelta(seconds=float(cell.value))
        if cell.kind == "datetime":
            return datetime.fromisoformat(cell.value)
        if cell.kind == "date":
            return date.fromisoformat(cell.value)
        return time.fromisoformat(cell.value)
    return value


def _upload(session: Session, project: Project, upload_id: uuid.UUID) -> CustomerUpload:
    result = session.exec(
        select(CustomerUpload).where(
            CustomerUpload.id == upload_id,
            CustomerUpload.project_id == project.id,
            CustomerUpload.tenant_id == project.tenant_id,
        )
    ).one_or_none()
    if result is None:
        raise LedgerError("ledger_upload_not_found", 404)
    return result


def _revision(
    session: Session, project: Project, revision_id: uuid.UUID
) -> CustomerLedgerRevision:
    result = session.exec(
        select(CustomerLedgerRevision).where(
            CustomerLedgerRevision.id == revision_id,
            CustomerLedgerRevision.project_id == project.id,
            CustomerLedgerRevision.tenant_id == project.tenant_id,
        )
    ).one_or_none()
    if result is None or not result.sealed:
        raise LedgerError("ledger_revision_not_found", 404)
    return result


def _source_rows(
    session: Session, project: Project, upload: CustomerUpload
) -> list[LedgerEntry]:
    artifact = session.exec(
        select(Artifact).where(
            Artifact.id == upload.artifact_id, Artifact.tenant_id == project.tenant_id
        )
    ).one_or_none()
    if artifact is None:
        raise LedgerError("ledger_source_invalid", 500)
    workbook = None
    try:
        path = uploads._artifact_path(
            artifact_root=settings.ARTIFACT_ROOT, artifact=artifact
        )
        if (
            path.stat().st_size > MAX_WORKBOOK_BYTES
            or path.stat().st_size != artifact.byte_size
        ):
            raise LedgerError("ledger_source_invalid", 500)
        with path.open("rb") as stream:
            digest = hashlib.file_digest(stream, "sha256").hexdigest()
        if digest != upload.raw_sha256 or digest != artifact.sha256:
            raise LedgerError("ledger_source_invalid", 500)
        validation = validate_customer_upload_workbook(path)
        if validation.record_count != upload.record_count:
            raise LedgerError("ledger_source_invalid", 500)
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=True)
        bounds = validation.bounds
        if bounds is None:
            raise LedgerError("ledger_source_invalid", 500)
        rows = workbook.worksheets[0].iter_rows(
            min_row=1,
            max_row=max(1, bounds.max_row),
            min_col=1,
            max_col=max(1, bounds.max_column, bounds.max_header_column),
            values_only=True,
        )
        headers = list(next(rows))
        indexes = {
            field: headers.index(header)
            for field, header in zip(FIELD_NAMES, HEADERS, strict=True)
            if header in headers
        }
        result: list[LedgerEntry] = []
        for source_row, row in enumerate(rows, start=2):
            if all(value is None for value in row):
                continue
            fields = {field: _cell(row[index]) for field, index in indexes.items()}
            result.append(
                LedgerEntry(
                    entry_id=uuid.uuid5(upload.id, f"row:{source_row}"),
                    position=len(result) + 1,
                    canonical_ip=normalize_ip(fields["asset_ip"]),
                    fields=cast(dict[FieldName, Cell], fields),
                )
            )
        if len(result) != upload.record_count:
            raise LedgerError("ledger_source_invalid", 500)
        return result
    except LedgerError:
        raise
    except Exception:
        raise LedgerError("ledger_source_invalid", 500) from None
    finally:
        if workbook is not None:
            workbook.close()


def load_rows(
    session: Session,
    project: Project,
    upload_id: uuid.UUID,
    revision_id: uuid.UUID | None,
) -> tuple[CustomerUpload, CustomerLedgerRevision | None, list[LedgerEntry]]:
    upload = _upload(session, project, upload_id)
    if revision_id is None:
        return upload, None, _source_rows(session, project, upload)
    revision = _revision(session, project, revision_id)
    if revision.upload_id != upload.id:
        raise LedgerError("ledger_scope_mismatch", 404)
    rows = session.exec(
        select(CustomerLedgerEntryVersion)
        .where(CustomerLedgerEntryVersion.revision_id == revision.id)
        .order_by(col(CustomerLedgerEntryVersion.position))
    ).all()
    if not rows:
        raise LedgerError("ledger_revision_invalid", 500)
    return (
        upload,
        revision,
        [LedgerEntry.model_validate(row, from_attributes=True) for row in rows],
    )


def read_page(
    session: Session,
    project: Project,
    *,
    upload_id: uuid.UUID | None,
    revision_id: uuid.UUID | None,
    original: bool,
    query: str,
    ip: str | None,
    archived: bool,
    skip: int,
    limit: int,
    can_write: bool,
) -> LedgerPage:
    current = project.current_customer_ledger_revision_id
    if revision_id is not None:
        upload_id = _revision(session, project, revision_id).upload_id
    elif upload_id is None:
        upload_id = project.current_customer_upload_id
        if not original:
            revision_id = current
    page = LedgerPage(
        project_id=project.id,
        upload_id=upload_id,
        revision_id=revision_id,
        current_upload_id=project.current_customer_upload_id,
        current_revision_id=current,
        current_profile_id=project.current_customer_upload_profile_id,
    )
    if upload_id is None:
        return page
    upload, revision, rows = load_rows(session, project, upload_id, revision_id)
    page.upload_sha256 = upload.raw_sha256
    page.filename = upload.display_filename
    page.source_created_at = upload.created_at
    page.revision = public_revision(revision) if revision else None
    page.can_edit = (
        can_write
        and upload_id == project.current_customer_upload_id
        and revision_id == current
    )
    # ponytail: bounded workbook/revision scan; indexed entry filtering if larger ledgers need it.
    active = [row for row in rows if not row.archived]
    page.total_records = len(active)
    page.unique_ips = len({row.canonical_ip for row in active})
    normalized_ip = normalize_ip(ip) if ip else None
    needle = query.casefold().strip()
    matches = [
        row
        for row in rows
        if row.archived == archived
        and (normalized_ip is None or row.canonical_ip == normalized_ip)
        and (
            not needle
            or needle
            in json.dumps(row.model_dump(mode="json"), ensure_ascii=False).casefold()
        )
    ]
    page.count = len(matches)
    page.data = matches[skip : skip + limit]
    return page


def recover(
    session: Session, project: Project, actor_id: uuid.UUID, key: str
) -> CustomerLedgerRevision | None:
    return session.exec(
        select(CustomerLedgerRevision).where(
            CustomerLedgerRevision.project_id == project.id,
            CustomerLedgerRevision.tenant_id == project.tenant_id,
            CustomerLedgerRevision.created_by == actor_id,
            CustomerLedgerRevision.operation_key == key,
        )
    ).one_or_none()


def _write_workbook(rows: list[LedgerEntry]) -> Path:
    root = settings.ARTIFACT_ROOT
    root.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=root, suffix=".xlsx", delete=False) as output:
        path = Path(output.name)
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.append(list(HEADERS))
        for row in rows:
            if row.archived:
                continue
            values = [_excel(row.fields.get(field)) for field in FIELD_NAMES]
            _validate_required_row(
                values,
                dict(zip(FIELD_NAMES, range(len(FIELD_NAMES)), strict=True)),
                row.position + 1,
            )
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = (
                        "s"  # User text is always a literal, never an Excel formula.
                    )
        workbook.save(path)
        validate_customer_upload_workbook(path)
        return path
    except Exception as error:
        path.unlink(missing_ok=True)
        if isinstance(
            error,
            (
                IllegalCharacterError,
                TypeError,
                ValueError,
                OverflowError,
            ),
        ):
            raise LedgerError("ledger_fields_invalid", 422) from None
        raise
    finally:
        workbook.close()


def save(
    session: Session,
    project: Project,
    *,
    actor_id: uuid.UUID,
    key: str,
    edit: LedgerEdit,
    ip_address: str | None,
) -> CustomerLedgerRevision:
    digest = hashlib.sha256(
        json.dumps(
            edit.model_dump(mode="json"),
            sort_keys=True,
            separators=(",", ":"),
            ensure_ascii=False,
        ).encode()
    ).hexdigest()
    existing = recover(session, project, actor_id, key)
    if existing is not None:
        if existing.request_sha256 != digest:
            raise LedgerError("ledger_key_conflict")
        return existing
    if (
        project.current_customer_upload_id != edit.expected_upload_id
        or project.current_customer_ledger_revision_id != edit.expected_revision_id
        or project.current_customer_upload_profile_id != edit.expected_profile_id
    ):
        raise LedgerError("ledger_version_conflict")
    base, _parent, rows = load_rows(
        session, project, edit.expected_upload_id, edit.expected_revision_id
    )
    if base.profile_id != edit.expected_profile_id:
        raise LedgerError("ledger_profile_changed")
    rows = copy.deepcopy(rows)
    if not edit.reason.strip():
        raise LedgerError("ledger_reason_required", 422)
    row = next((item for item in rows if item.entry_id == edit.entry_id), None)
    previous_fields = copy.deepcopy(row.fields) if row is not None else None
    if edit.operation == "add":
        if edit.entry_id is not None or not edit.fields:
            raise LedgerError("ledger_operation_invalid", 422)
        row = LedgerEntry(
            entry_id=uuid.uuid4(),
            position=len(rows) + 1,
            canonical_ip="",
            fields=edit.fields,
        )
        rows.append(row)
    elif row is None or row.archived:
        raise LedgerError("ledger_entry_not_found", 404)
    if edit.operation in {"update", "add"}:
        if not edit.fields or edit.management is not None:
            raise LedgerError("ledger_operation_invalid", 422)
        row.fields.update(edit.fields)
    elif edit.operation == "archive":
        if edit.fields or edit.management is not None:
            raise LedgerError("ledger_operation_invalid", 422)
        row.archived = True
    else:
        if edit.fields or edit.management is None:
            raise LedgerError("ledger_operation_invalid", 422)
        row.management = edit.management
    input_changed = edit.operation in {"add", "archive"} or (
        edit.operation == "update" and row.fields != previous_fields
    )
    path = None
    promoted: list[Path] = []
    committing = False
    try:
        if input_changed:
            if not any(not item.archived for item in rows):
                raise LedgerError("ledger_last_entry", 422)
            for item in rows:
                if not item.archived:
                    item.canonical_ip = normalize_ip(
                        str(item.fields.get("asset_ip", ""))
                    )
            path = _write_workbook(rows)
            size = path.stat().st_size
            with path.open("rb") as stream:
                raw_hash = hashlib.file_digest(stream, "sha256").hexdigest()
            (settings.ARTIFACT_ROOT / "customer_uploads").mkdir(exist_ok=True)
            upload, _created = uploads.accept_customer_upload(
                session=session,
                project=project,
                streamed_upload=uploads.StreamedCustomerUpload(
                    path, f"manual-derived-{uuid.uuid4()}.xlsx", size, raw_hash
                ),
                artifact_root=settings.ARTIFACT_ROOT,
                actor_subject=str(actor_id),
                ip_address=ip_address,
                commit=False,
                promoted_paths=promoted,
            )
        else:
            upload = base
        revision = CustomerLedgerRevision(
            tenant_id=project.tenant_id,
            project_id=project.id,
            parent_revision_id=edit.expected_revision_id,
            base_upload_id=base.id,
            upload_id=upload.id,
            created_by=actor_id,
            operation_key=key,
            request_sha256=digest,
            reason=edit.reason.strip(),
            input_changed=input_changed,
        )
        session.add(revision)
        session.flush()
        session.add_all(
            [
                CustomerLedgerEntryVersion(revision_id=revision.id, **item.model_dump())
                for item in rows
            ]
        )
        session.flush()
        before = {
            "upload_id": str(project.current_customer_upload_id),
            "revision_id": str(project.current_customer_ledger_revision_id)
            if project.current_customer_ledger_revision_id
            else None,
        }
        project.current_customer_upload_id = upload.id
        project.current_customer_ledger_revision_id = revision.id
        session.add(project)
        session.add(
            AuditEvent(
                tenant_id=project.tenant_id,
                project_id=project.id,
                actor_subject=str(actor_id),
                actor_type="user",
                action="customer_ledger.saved",
                target_type="customer_ledger_revision",
                target_id=revision.id,
                before_data=before,
                after_data={
                    "upload_id": str(upload.id),
                    "revision_id": str(revision.id),
                    "input_changed": input_changed,
                    "record_count": sum(not item.archived for item in rows),
                },
                ip_address=ip_address,
            )
        )
        session.flush()
        revision.sealed = True
        session.add(revision)
        session.flush()
        session.expunge(revision)
        committing = True
        session.commit()
        return revision
    except Exception as error:
        session.rollback()
        if committing:
            try:
                with Session(session.get_bind()) as check:
                    committed = recover(check, project, actor_id, key)
                    if committed is not None:
                        check.expunge(committed)
                        return committed
            except SQLAlchemyError:
                # Preserve potentially committed bytes; recovery uses the same operation key.
                raise LedgerError("ledger_commit_unknown", 503) from None
        for promoted_path in promoted:
            promoted_path.unlink(missing_ok=True)
        if isinstance(error, LedgerError):
            raise
        if isinstance(
            error,
            CustomerUploadValidationError
            | IPRecordContractError
            | ValueError
            | OverflowError,
        ):
            raise LedgerError("ledger_fields_invalid", 422) from None
        if isinstance(error, uploads.CustomerUploadAcceptanceError):
            raise LedgerError(
                error.code, 500 if error.code == "upload_storage_failed" else 422
            ) from None
        if isinstance(error, SQLAlchemyError | OSError):
            raise LedgerError("ledger_save_failed", 500) from None
        raise
    finally:
        if path is not None:
            path.unlink(missing_ok=True)
